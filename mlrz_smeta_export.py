"""
Смета — АО «Московский ЛРЗ» / ООО «СпецЭкоСервис», Договор №2025.223692.
Срок 01.10.2025-30.09.2026 (12 мес). Цена — РЕАЛЬНАЯ, из Приложения №1
(поединичные расценки за м², сумма сошлась с ценой договора до копейки).

Смета не 'предсказывает' цену — она уже известна. Структура здесь: срок->
цена/мес->цена/год->численность (два метода, честно оба)->материалы
(резолвер, наша база)->разбивка по видам помещений.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from knowledge_base import ObjectParams, estimate_consumable_quantity, load_consumable_norms
from models import TZItem
from price_parser import load_almin_price_list, load_general_opt_price_list
from resolver import resolve_all
from supplier_fallback import stub_supplier_adapter
from unit_price import estimate_purchase_quantity

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FONT = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="666666")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY_FMT = '#,##0.00 ₽;(#,##0.00 ₽);-'

CONTRACT_MONTHS = 12
VAT_RATE = 0.05
AREAS = [
    ("Административные помещения", 8764.0, 14.67),
    ("Производственные помещения", 79460.0, 14.89),
    ("Складские помещения", 2478.0, 14.80),
    ("Территория", 15225.0, 7.31),
]
AREA_INDOOR = 8764.0 + 79460.0 + 2478.0
EMISS_SALARY = 79181.10

wb = Workbook()
ws = wb.active
ws.title = "Смета"
ws.sheet_view.showGridLines = False
for col, width in zip("ABCDEFG", [58, 22, 18, 16, 16, 16, 14]):
    ws.column_dimensions[col].width = width

row = 1
ws.cell(row=row, column=1, value="Смета — АО «Московский ЛРЗ» (Договор №2025.223692)").font = Font(
    name=FONT_NAME, bold=True, size=14
)
row += 1
ws.cell(row=row, column=1, value=f"Дата: {date.today().isoformat()} | Цена — РЕАЛЬНАЯ из Приложения №1, не оценка").font = NOTE_FONT
row += 2
ws.cell(row=row, column=1, value="Жёлтые ячейки — допущения").fill = INPUT_FILL
ws.cell(row=row, column=1).font = Font(name=FONT_NAME, size=9)
row += 2

# ============================== 1. Срок ==============================
ws.cell(row=row, column=1, value="1. Срок контракта").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="Срок оказания услуг, мес (01.10.2025-30.09.2026)")
r_months = row
ws.cell(row=row, column=2, value=CONTRACT_MONTHS)
row += 1
ws.cell(row=row, column=1, value="Площадь всего, м² (помещения + территория)")
ws.cell(row=row, column=2, value=AREA_INDOOR + 15225.0)
row += 2
for r in range(row - 3, row - 1):
    for c in (1, 2):
        ws.cell(row=r, column=c).border = BORDER

# ============================== 2. Цена (реальная, по видам) ==============================
ws.cell(row=row, column=1, value="2. Цена — реальные поединичные расценки (Приложение №1)").font = SECTION_FONT
row += 1
headers = ["Вид помещения", "Площадь, м²", "Ставка ₽/м²/мес (без НДС)", "Сумма/мес, ₽"]
for c, h in enumerate(headers, start=1):
    ws.cell(row=row, column=c, value=h)
    ws.cell(row=row, column=c).fill = HEADER_FILL
    ws.cell(row=row, column=c).font = HEADER_FONT
    ws.cell(row=row, column=c).border = BORDER
row += 1
area_first_row = row
for name, area, rate in AREAS:
    ws.cell(row=row, column=1, value=name)
    ws.cell(row=row, column=2, value=area)
    ws.cell(row=row, column=3, value=rate).number_format = CURRENCY_FMT
    ws.cell(row=row, column=4, value=f"=B{row}*C{row}").number_format = CURRENCY_FMT
    for c in range(1, 5):
        ws.cell(row=row, column=c).border = BORDER
    row += 1
area_last_row = row - 1

ws.cell(row=row, column=1, value="ЦЕНА В МЕСЯЦ (без НДС)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
r_price_month_novat = row
ws.cell(row=row, column=4, value=f"=SUM(D{area_first_row}:D{area_last_row})").number_format = CURRENCY_FMT
ws.cell(row=row, column=4).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="ЦЕНА В МЕСЯЦ (с НДС 5%)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
r_price_month = row
ws.cell(row=row, column=4, value=f"=D{r_price_month_novat}*{1+VAT_RATE}").number_format = CURRENCY_FMT
ws.cell(row=row, column=4).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="ЦЕНА В ГОД (с НДС)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
ws.cell(row=row, column=4, value=f"=D{r_price_month}*{CONTRACT_MONTHS}").number_format = CURRENCY_FMT
ws.cell(row=row, column=4).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 2

# ============================== 3. Численность ==============================
ws.cell(row=row, column=1, value="3. Численность и ФОТ — РЕАЛЬНЫЙ табель объекта").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="ФОТ, ₽/мес (сумма по реальному табелю — 18 позиций, разные оклады по сменам)")
r_fot_real = row
ws.cell(row=row, column=2, value=870000.00).number_format = CURRENCY_FMT
row += 1
ws.cell(row=row, column=1, value="ФОТ-надбавка 30,2% (обязательна)")
r_fot_pct = row
ws.cell(row=row, column=2, value=f"=B{r_fot_real}*0.302").number_format = CURRENCY_FMT
row += 1
ws.cell(row=row, column=1, value="Справочно: наша прежняя оценка (11 чел. × ЕМИСС) была 870 992 ₽ — разошлась всего на 0,11%, "
                                    "хотя реальная структура другая (18 позиций по сменам, не 11 человек по одной ставке).").font = NOTE_FONT
row += 2
for r in range(row - 4, row - 1):
    for c in (1, 2):
        ws.cell(row=r, column=c).border = BORDER

# ============================== 4. Материалы ==============================
ws.cell(row=row, column=1, value="4. Расходные материалы (наша база, справочно — не входит в цену выше, она уже финальна)").font = SECTION_FONT
row += 1
params = ObjectParams(area_sqm=AREA_INDOOR, cleaning_days=26, general_days=1, staff_count=11, contract_months=1)
norms = load_consumable_norms()
LARGE_OBJECT_CORRECTION = 1 / 7
COUNT_BASED = {"мусорный_мешок_крупный", "салфетки_универсальные"}
DURABLE = {"комплект_уборки"}
_price_lists = load_almin_price_list(discount_pct=5.0) + load_general_opt_price_list(supplier="ТК Сервис")

_tz_items, _norm_qtys = [], []
for norm_id in norms:
    if norm_id == "моющее_средство_полы_стены" or norm_id in DURABLE:
        continue
    qty, unit, _ = estimate_consumable_quantity(norm_id, params)
    if norm_id not in COUNT_BASED and norm_id != "моющее_средство_полы_механизированная_уборка":
        qty = round(qty * LARGE_OBJECT_CORRECTION, 2)
    _tz_items.append(TZItem(raw_name=norms[norm_id]["display_name"], qty=qty, unit=unit))
    _norm_qtys.append((qty, unit))
_match_results = resolve_all(_tz_items, _price_lists, supplier_adapters={"almin": stub_supplier_adapter})

headers_mat = ["Материал", "Найдено в прайсе", "Поставщик", "Упаковок", "Цена/уп, ₽", "Сумма/мес, ₽"]
for c, h in enumerate(headers_mat, start=1):
    ws.cell(row=row, column=c, value=h)
    ws.cell(row=row, column=c).fill = HEADER_FILL
    ws.cell(row=row, column=c).font = HEADER_FONT
    ws.cell(row=row, column=c).border = BORDER
row += 1
mat_first_row = row
for r, (qty, unit) in zip(_match_results, _norm_qtys):
    ws.cell(row=row, column=1, value=f"{r.tz_item.raw_name[:38]} ({qty} {unit})")
    matched_name = r.matched_item.name[:26] if r.matched_item else "НЕ НАЙДЕНО"
    ws.cell(row=row, column=2, value=f"{matched_name} (c{r.confidence:.2f})" if r.matched_item else matched_name)
    ws.cell(row=row, column=3, value=r.matched_item.supplier if r.matched_item else "—")
    packs = estimate_purchase_quantity(qty, unit, r.matched_item) if r.matched_item and unit in ("л", "кг", "шт") else None
    packs_cell = ws.cell(row=row, column=4, value=packs if packs is not None else qty)
    if packs is None:
        packs_cell.fill = INPUT_FILL
    price = r.matched_item.final_price if r.matched_item else 0
    price_cell = ws.cell(row=row, column=5, value=price)
    price_cell.number_format = CURRENCY_FMT
    if not r.matched_item:
        price_cell.fill = INPUT_FILL
    ws.cell(row=row, column=6, value=f"=D{row}*E{row}").number_format = CURRENCY_FMT
    for c in range(1, 7):
        ws.cell(row=row, column=c).border = BORDER
        if c != 1:
            ws.cell(row=row, column=c).font = Font(name=FONT_NAME, size=9)
    row += 1
mat_last_row = row - 1
ws.cell(row=row, column=1, value="ИТОГО материалы/мес (наша модель, сверена по ₽/м² с реальными данными Краун)").font = Font(name=FONT_NAME, bold=True)
r_mat = row
ws.cell(row=row, column=6, value=f"=SUM(F{mat_first_row}:F{mat_last_row})").number_format = CURRENCY_FMT
ws.cell(row=row, column=6).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value="Площадь для норм взята только по помещениям (90 702 м²) — территория считается отдельно, здесь не включена.").font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
row += 2

ws.cell(row=row, column=1, value="5. Сверка себестоимости с реальной ценой").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="ИТОГО себестоимость (реальный ФОТ+надбавка+наши материалы)").font = TOTAL_FONT
r_total_cost = row
ws.cell(row=row, column=2, value=f"=B{r_fot_real}+B{r_fot_pct}+F{r_mat}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = TOTAL_FONT
row += 1
ws.cell(row=row, column=1, value="Реальная цена (без НДС)")
ws.cell(row=row, column=2, value=f"=D{r_price_month_novat}").number_format = CURRENCY_FMT
row += 1
ws.cell(row=row, column=1, value="Маржа (накладные+прибыль)").font = Font(name=FONT_NAME, bold=True)
ws.cell(row=row, column=2, value=f"=D{r_price_month_novat}-B{r_total_cost}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value=f"=(D{r_price_month_novat}-B{r_total_cost})/D{r_price_month_novat}")
ws.cell(row=row, column=1, value="Маржа, %")
ws.cell(row=row, column=2, value=f"=(D{r_price_month_novat}-B{r_total_cost})/D{r_price_month_novat}").number_format = "0.0%"
row += 1
ws.cell(row=row, column=1,
        value="ФОТ теперь реальный (подтверждён табелем, расхождение с прежней оценкой — 0,11%). Маржа получилась "
              "тонкой (~2%) — проверка по ₽/м² (сверка с реальными данными Краун, 2.50-4.17 руб/м²/мес) показала, "
              "что наша модель материалов НЕ завышена, попадает в реальный диапазон. Похоже, тонкая маржа отражает "
              "реальную структуру этого конкретного контракта (конкурентная закупка), а не ошибку модели. Есть "
              "также реальный счёт на материалы по одной зоне доставки объекта (40 514 руб) — но масштаб этой зоны "
              "относительно всего завода неизвестен, использовать для уточнения общей оценки пока нельзя.").font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)

wb.save("smeta_mlrz.xlsx")
print("Смета сохранена: smeta_mlrz.xlsx")
