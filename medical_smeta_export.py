"""
Смета по ТЗ "Оказание услуг по уборке помещений в 2027 году" — медицинская
организация, 3 корпуса (Ермолаевский пер./Даев пер./Протопоповский пер.),
16 026,9 м², срок 01.01.2027-31.12.2027 (12 мес).

ЧЕСТНО: в отличие от Лужников, здесь НЕТ приложенной официальной сметы —
это первый проход по методике (как самая первая смета, Большой Головин),
не сверенный с government-audited цифрами. Медицинская специфика (до 4
уборок в день по регламенту, класс чистоты Г/В) заметно интенсивнее
обычного офиса — ставьте под сомнение параметры ниже сильнее обычного.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FONT = Font(name=FONT_NAME, bold=True, size=12)
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="666666")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY_FMT = '#,##0.00 ₽;(#,##0.00 ₽);-'

AREA_SQM = 16026.9
CONTRACT_MONTHS = 12
EMISS_SALARY = 79181.10  # тот же официальный ориентир ЕМИСС, что и для Лужников

BUILDINGS = [
    ("Ермолаевский пер., д.22/26, стр.1", 8, 5454.5),
    ("Даев пер., д.3, стр.1", 7, 7629.8),
    ("Протопоповский пер., д.19, стр.15", 5, 2942.6),
]

wb = Workbook()
ws = wb.active
ws.title = "Смета"
ws.sheet_view.showGridLines = False
for col, width in zip("ABCDE", [55, 22, 16, 16, 45]):
    ws.column_dimensions[col].width = width

row = 1
ws.cell(row=row, column=1, value="Смета — Уборка помещений медицинской организации, 2027 год").font = Font(
    name=FONT_NAME, bold=True, size=14
)
row += 1
ws.cell(row=row, column=1, value=f"Дата: {date.today().isoformat()}").font = Font(name=FONT_NAME, size=10)
row += 2

ws.cell(row=row, column=1, value="ВАЖНО: официальной сметы заказчика к этому ТЗ не приложено — это первый "
                                    "проход по методике, не сверенный с government-audited цифрами.").font = Font(
    name=FONT_NAME, italic=True, size=9, color="C00000"
)
row += 2

ws.cell(row=row, column=1, value="Жёлтые ячейки — допущения, требуют проверки эксперта").fill = INPUT_FILL
ws.cell(row=row, column=1).font = Font(name=FONT_NAME, size=9)
row += 2

# ============================== 0. Объекты ==============================
ws.cell(row=row, column=1, value="0. Объекты (из Приложения 3, реальные площади)").font = SECTION_FONT
row += 1
headers0 = ["Корпус", "Этажей", "Площадь, м²"]
for c, h in enumerate(headers0, start=1):
    ws.cell(row=row, column=c, value=h)
    ws.cell(row=row, column=c).fill = HEADER_FILL
    ws.cell(row=row, column=c).font = HEADER_FONT
    ws.cell(row=row, column=c).border = BORDER
row += 1
b_first_row = row
for name, floors, area in BUILDINGS:
    ws.cell(row=row, column=1, value=name)
    ws.cell(row=row, column=2, value=floors)
    ws.cell(row=row, column=3, value=area).number_format = CURRENCY_FMT.replace(" ₽", "")
    for c in range(1, 4):
        ws.cell(row=row, column=c).border = BORDER
    row += 1
b_last_row = row - 1
ws.cell(row=row, column=1, value="ИТОГО площадь").font = TOTAL_FONT
r_area_total = row
ws.cell(row=row, column=3, value=f"=SUM(C{b_first_row}:C{b_last_row})")
ws.cell(row=row, column=3).font = TOTAL_FONT
row += 2

# ============================== 1. Трудозатраты ==============================
ws.cell(row=row, column=1, value="1. Трудозатраты").font = SECTION_FONT
row += 1

def data_row(label, value, editable=False, fmt=None, note=None):
    global row
    ws.cell(row=row, column=1, value=label).font = Font(name=FONT_NAME, size=10)
    cell = ws.cell(row=row, column=2, value=value)
    cell.font = Font(name=FONT_NAME, size=10)
    cell.border = BORDER
    if editable:
        cell.fill = INPUT_FILL
    if fmt:
        cell.number_format = fmt
    if note:
        ws.cell(row=row, column=5, value=note).font = NOTE_FONT
    row += 1
    return row - 1

r_norm = data_row("Норматив, м²/чел (базовый, с поправкой на многократную уборку)", 336.7, editable=True,
                    note="370.4 (Роструд) / 1.1 (поправка за >1 уборки/смену). Медицина по регламенту требует "
                         "до 4 обработок/день по отдельным элементам — возможно, стоит взять норматив ЕЩЁ ниже.")
row += 1
ws.cell(row=row, column=1, value="Численность (расчёт)").font = Font(name=FONT_NAME, bold=True)
r_staff = row
ws.cell(row=row, column=2, value=f"=ROUNDUP(C{r_area_total}/B{r_norm},0)")
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1

r_salary = data_row("Оклад, ₽/мес (ЕМИСС/Росстат, комплексное обслуживание помещений, Москва 2025)",
                      EMISS_SALARY, editable=True, fmt=CURRENCY_FMT)
r_months = data_row("Срок контракта, мес", CONTRACT_MONTHS, editable=True)
row += 1

ws.cell(row=row, column=1, value="ИТОГО ФОТ за контракт").font = TOTAL_FONT
r_fot_total = row
ws.cell(row=row, column=2, value=f"=B{r_staff}*B{r_salary}*B{r_months}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = TOTAL_FONT
row += 2

# ============================== 2. Материалы (грубая оценка) ==============================
ws.cell(row=row, column=1, value="2. Материалы (предварительная оценка по площади)").font = SECTION_FONT
row += 1
r_mat_norm = data_row("Норма расхода, ₽/м²/год (медицинские нормы дезинфекции — ВЫШЕ обычного офиса)",
                        450, editable=True,
                        note="ГРУБАЯ оценка агента — нет официальных расценок для этого объекта (в отличие от "
                             "Лужников). Медицинская дезинфекция (СанПиН, классы Г/В) обычно дороже обычного "
                             "клининга. ПРОВЕРИТЬ по факту закупочных цен.")
row += 1
ws.cell(row=row, column=1, value="ИТОГО материалы за контракт").font = TOTAL_FONT
r_mat_total = row
ws.cell(row=row, column=2, value=f"=C{r_area_total}*B{r_mat_norm}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = TOTAL_FONT
row += 2

# ============================== Итог ==============================
ws.cell(row=row, column=1, value="ИТОГО СЕБЕСТОИМОСТЬ (ФОТ + материалы)").font = Font(
    name=FONT_NAME, bold=True, size=13, color="C00000"
)
r_cost = row
ws.cell(row=row, column=2, value=f"=B{r_fot_total}+B{r_mat_total}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="Себестоимость в месяц").font = Font(name=FONT_NAME, bold=True)
ws.cell(row=row, column=2, value=f"=B{r_cost}/B{r_months}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 2

ws.cell(row=row, column=1,
        value="НЕ включено: накладные расходы, сметная прибыль, НДС — добавьте по вашей практике (см. "
              "примеры Лужники/Большой Головин). Периодичность по регламенту (до 4х уборок/день для отдельных "
              "элементов, классы Г/В) не разложена по операциям — норматив выше УСРЕДНЁННЫЙ, не детальный ГЭСН-расчёт.")
ws.cell(row=row, column=1).font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)

wb.save("smeta_medical_2027.xlsx")
print("Смета сохранена: smeta_medical_2027.xlsx")
