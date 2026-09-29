"""
Смета — АО «Московский ЛРЗ» / ООО «СпецЭкоСервис», Договор №2025.223692.
Срок 01.10.2025-30.09.2026 (12 мес). Цена — РЕАЛЬНАЯ, из Приложения №1
(поединичные расценки за м², сумма сошлась с ценой договора до копейки).

v2: та же смета, что и mlrz_smeta_export.py, но собрана через общий
smeta_builder.py (систематизация скриптов смет) — секция материалов
теперь с настоящими Excel-формулами (норматив×площадь×дни, а не
захардкоженное число), чтобы проверить, что рефакторинг даёт тот же
результат, что и оригинал, и что recalc.py не находит ошибок.
"""
from datetime import date

from openpyxl import Workbook

from knowledge_base import ObjectParams
from price_parser import load_almin_price_list, load_general_opt_price_list
from smeta_builder import (
    CURRENCY_FMT, FINAL_FONT, NOTE_FONT, SECTION_FONT, TOTAL_FONT,
    SmetaSheet, build_contract_term_section, build_materials_table,
)
from supplier_fallback import stub_supplier_adapter

CONTRACT_MONTHS = 12
VAT_RATE = 0.05
AREAS = [
    ("Административные помещения", 8764.0, 14.67),
    ("Производственные помещения", 79460.0, 14.89),
    ("Складские помещения", 2478.0, 14.80),
    ("Территория", 15225.0, 7.31),
]
AREA_INDOOR = 8764.0 + 79460.0 + 2478.0

wb = Workbook()
sheet = SmetaSheet.new(wb, "Смета", {c: w for c, w in zip("ABCDEFGHI", [46, 16, 16, 8, 34, 22, 12, 14, 14])})
sheet.title_block(
    "АО «Московский ЛРЗ» (Договор №2025.223692)",
    subtitle="Цена — РЕАЛЬНАЯ из Приложения №1, не оценка. Сгенерировано через smeta_builder.py (v2).",
)

# ============================== 1. Срок ==============================
term_refs = build_contract_term_section(sheet, CONTRACT_MONTHS, area_total_sqm=AREA_INDOOR + 15225.0)

# ============================== 2. Цена — реальные поединичные расценки ==============================
sheet.section("2. Цена — реальные поединичные расценки (Приложение №1)")
headers = ["Вид помещения", "Площадь, м²", "Ставка ₽/м²/мес (без НДС)", "Сумма/мес, ₽"]
sheet.table_header(headers)
area_first_row = sheet.row
for name, area, rate in AREAS:
    r = sheet.row
    sheet.ws.cell(row=r, column=1, value=name)
    sheet.ws.cell(row=r, column=2, value=area)
    sheet.ws.cell(row=r, column=3, value=rate).number_format = CURRENCY_FMT
    sheet.ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = CURRENCY_FMT
    sheet.row += 1
area_last_row = sheet.row - 1

r_price_month_novat = sheet.row
sheet.ws.cell(row=r_price_month_novat, column=1, value="ЦЕНА В МЕСЯЦ (без НДС)").font = FINAL_FONT
sheet.ws.cell(row=r_price_month_novat, column=4,
              value=f"=SUM(D{area_first_row}:D{area_last_row})").number_format = CURRENCY_FMT
sheet.ws.cell(row=r_price_month_novat, column=4).font = FINAL_FONT
sheet.row += 1
r_price_month = sheet.row
sheet.ws.cell(row=r_price_month, column=1, value=f"ЦЕНА В МЕСЯЦ (с НДС {VAT_RATE*100:.0f}%)").font = FINAL_FONT
sheet.ws.cell(row=r_price_month, column=4, value=f"=D{r_price_month_novat}*{1+VAT_RATE}").number_format = CURRENCY_FMT
sheet.ws.cell(row=r_price_month, column=4).font = FINAL_FONT
sheet.row += 1
sheet.ws.cell(row=sheet.row, column=1, value="ЦЕНА В ГОД (с НДС)").font = FINAL_FONT
sheet.ws.cell(row=sheet.row, column=4, value=f"=D{r_price_month}*{CONTRACT_MONTHS}").number_format = CURRENCY_FMT
sheet.ws.cell(row=sheet.row, column=4).font = FINAL_FONT
sheet.row += 2

# ============================== 3. Численность — реальный табель ==============================
sheet.section("3. Численность и ФОТ — РЕАЛЬНЫЙ табель объекта")
r_fot_real = sheet.value_row("ФОТ, ₽/мес (сумма по реальному табелю — 18 позиций, разные оклады по сменам)",
                              870000.00, fmt=CURRENCY_FMT)
r_fot_pct = sheet.row
sheet.ws.cell(row=r_fot_pct, column=1, value="ФОТ-надбавка 30,2% (обязательна)")
sheet.ws.cell(row=r_fot_pct, column=2, value=f"=B{r_fot_real}*0.302").number_format = CURRENCY_FMT
sheet.row += 1
sheet.note("Справочно: наша прежняя оценка (11 чел. × ЕМИСС) была 870 992 ₽ — разошлась всего на 0,11%, "
           "хотя реальная структура другая (18 позиций по сменам, не 11 человек по одной ставке).")
sheet.blank()

# ============================== 4. Материалы — через build_materials_table ==============================
params = ObjectParams(area_sqm=AREA_INDOOR, cleaning_days=26, general_days=1, staff_count=11, contract_months=1)
price_lists = load_almin_price_list(discount_pct=5.0) + load_general_opt_price_list(supplier="ТК Сервис")

mat = build_materials_table(
    sheet, params, price_lists,
    supplier_adapters={"almin": stub_supplier_adapter},
    exclude_norms={"моющее_средство_полы_стены", "комплект_уборки"},
    no_correction_norms={"мусорный_мешок_крупный", "салфетки_универсальные",
                          "моющее_средство_полы_механизированная_уборка"},
    correction_factor=1 / 7,
    correction_note="Поправка ×1/7 — крупный объект с механизированной/дозирующей уборкой, норма калибрована "
                     "на малый офис (см. knowledge_base/consumable_norms_scale_limitation.json). Площадь для норм "
                     "взята только по помещениям (90 702 м²) — территория считается отдельно, здесь не включена.",
)

sheet.section("5. Сверка себестоимости с реальной ценой")
r_total_cost = sheet.row
sheet.ws.cell(row=r_total_cost, column=1, value="ИТОГО себестоимость (реальный ФОТ+надбавка+наши материалы)").font = TOTAL_FONT
sheet.ws.cell(row=r_total_cost, column=2,
              value=f"=B{r_fot_real}+B{r_fot_pct}+{mat.total_cell}").number_format = CURRENCY_FMT
sheet.ws.cell(row=r_total_cost, column=2).font = TOTAL_FONT
sheet.row += 1
sheet.value_row("Реальная цена (без НДС)", f"=D{r_price_month_novat}", fmt=CURRENCY_FMT)
r_margin = sheet.row
sheet.ws.cell(row=r_margin, column=1, value="Маржа (накладные+прибыль)").font = TOTAL_FONT
sheet.ws.cell(row=r_margin, column=2, value=f"=D{r_price_month_novat}-B{r_total_cost}").number_format = CURRENCY_FMT
sheet.ws.cell(row=r_margin, column=2).font = TOTAL_FONT
sheet.row += 1
sheet.value_row("Маржа, %", f"=(D{r_price_month_novat}-B{r_total_cost})/D{r_price_month_novat}", fmt="0.0%")
sheet.note(
    "ФОТ реальный (подтверждён табелем). Проверка по ₽/м² (сверка с реальными данными Краун, 2.50-4.17 руб/м²/мес) "
    "показала, что наша модель материалов НЕ завышена. v2: смета собрана через smeta_builder.py — количество "
    "материалов теперь настоящая Excel-формула (норматив×площадь×дни×коэффициент коррекции), эксперт может "
    "поменять любую из этих ячеек и увидеть пересчёт, а не только цену за упаковку."
)

if mat.unresolved:
    sheet.note(f"Не найдено в прайсе (нужен ручной ввод цены): {', '.join(mat.unresolved)}")

wb.save("smeta_mlrz_v2.xlsx")
print("Смета сохранена: smeta_mlrz_v2.xlsx")
print(f"Материалы: строки {mat.first_row}-{mat.last_row}, итого в {mat.total_cell}, не найдено: {mat.unresolved}")
