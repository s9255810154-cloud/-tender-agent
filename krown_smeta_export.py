"""
Смета — Договор ООО «Краун» / ООО «СпецЭкоСервис», Озерковская наб., 28
стр.3. В отличие от прошлых смет, здесь есть РЕАЛЬНАЯ согласованная цена
и РЕАЛЬНОЕ штатное расписание по ролям (не оценка) — задача не "предложить
ставку", а проверить, сходится ли наша модель издержек с уже действующим
контрактом.

Структура — по стандарту (срок -> цена/мес -> цена/год -> численность ->
материалы с автосопоставлением и формулами -> остальное).
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from knowledge_base import ObjectParams, estimate_consumable_quantity, load_consumable_norms
from models import TZItem
from price_parser import load_almin_price_list, load_general_opt_price_list
from resolver import resolve_all
from supplier_fallback import stub_supplier_adapter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12)
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="666666")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY_FMT = '#,##0.00 ₽;(#,##0.00 ₽);-'

AREA_SQM = 15359.27
CONTRACT_MONTHS = 36
REAL_PRICE_WINTER = 2334071.96
REAL_PRICE_SUMMER = 2004000.00
REAL_PRICE_YEAR_1 = 25698359.80
REAL_PRICE_TOTAL_3Y = 80219999.96

STAFF_SUMMER = {
    "Уборщик производственных и служебных помещений": 15,
    "Бригадир": 1,
    "Уборщик территории": 1,
    "Оператор поломоечной машины": 1,
    "Автомойщик": 1,
    "Менеджер объекта": 1,
}
STAFF_WINTER_EXTRA = {
    "Уборщик территории": 1,
    "Оператор поломоечной машины": 1,
}
RATES = {
    "Уборщик производственных и служебных помещений": (80000, "проверено — salary_rules.json, Москва"),
    "Уборщик территории": (75000, "ОЦЕНКА — по аналогии с дворником (real_salary_benchmarks.json)"),
    "Автомойщик": (90000, "проверено — из этого же договора, salary_rules.json"),
    "Бригадир": (95000, "ОЦЕНКА — нет подтверждённой ставки"),
    "Оператор поломоечной машины": (85000, "ОЦЕНКА — нет подтверждённой ставки"),
    "Менеджер объекта": (110000, "ОЦЕНКА — нет подтверждённой ставки"),
}

wb = Workbook()
ws = wb.active
ws.title = "Смета"
ws.sheet_view.showGridLines = False
for col, width in zip("ABCDEFGH", [55, 30, 16, 14, 14, 16, 14, 14]):
    ws.column_dimensions[col].width = width

row = 1
ws.cell(row=row, column=1, value="Смета — Озерковская наб., 28 стр.3 (Краун / СпецЭкоСервис)").font = Font(
    name=FONT_NAME, bold=True, size=14
)
row += 1
ws.cell(row=row, column=1, value=f"Дата: {date.today().isoformat()} | Реальный подписанный договор — сверка модели, не предложение цены").font = Font(
    name=FONT_NAME, size=10, italic=True
)
row += 2
ws.cell(row=row, column=1, value="Жёлтые ячейки — допущения/оценки").fill = INPUT_FILL
ws.cell(row=row, column=1).font = Font(name=FONT_NAME, size=9)
row += 2

ws.cell(row=row, column=1, value="1. Срок контракта").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="Срок оказания услуг, мес (01.12.2025 - 30.11.2028)")
r_months = row
ws.cell(row=row, column=2, value=CONTRACT_MONTHS)
row += 1
ws.cell(row=row, column=1, value="Площадь помещений, м² (9-этажный офисный комплекс)")
r_area = row
ws.cell(row=row, column=2, value=AREA_SQM)
row += 2
for r in range(row - 3, row - 1):
    for c in (1, 2):
        ws.cell(row=r, column=c).border = BORDER

ws.cell(row=row, column=1, value="2. Цена по договору (реальная, согласованная)").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="ЦЕНА В МЕСЯЦ — зимний период (5 мес)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
r_price_winter = row
ws.cell(row=row, column=2, value=REAL_PRICE_WINTER).number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="ЦЕНА В МЕСЯЦ — летний период (7 мес)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
r_price_summer = row
ws.cell(row=row, column=2, value=REAL_PRICE_SUMMER).number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="ЦЕНА В ГОД (базовый год, с НДС 5%)").font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
ws.cell(row=row, column=2, value=REAL_PRICE_YEAR_1).number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True, size=13, color="C00000")
row += 1
ws.cell(row=row, column=1, value="Итого за весь договор (3 года, с индексацией ИПЦ 104%/год)")
ws.cell(row=row, column=2, value=REAL_PRICE_TOTAL_3Y).number_format = CURRENCY_FMT
row += 2
for r in range(row - 5, row - 1):
    for c in (1, 2):
        ws.cell(row=r, column=c).border = BORDER

ws.cell(row=row, column=1, value="3. Численность персонала (реальное штатное расписание из договора)").font = SECTION_FONT
row += 1
headers_staff = ["Должность", "Чел. (лето)", "Чел. доп. (зима)", "Ставка, ₽/мес", "Статус ставки"]
for c, h in enumerate(headers_staff, start=1):
    ws.cell(row=row, column=c, value=h)
    ws.cell(row=row, column=c).fill = HEADER_FILL
    ws.cell(row=row, column=c).font = HEADER_FONT
    ws.cell(row=row, column=c).border = BORDER
row += 1
staff_first_row = row
for role, summer_count in STAFF_SUMMER.items():
    winter_extra = STAFF_WINTER_EXTRA.get(role, 0)
    rate, rate_status = RATES[role]
    ws.cell(row=row, column=1, value=role)
    ws.cell(row=row, column=2, value=summer_count)
    ws.cell(row=row, column=3, value=winter_extra)
    rate_cell = ws.cell(row=row, column=4, value=rate)
    rate_cell.number_format = CURRENCY_FMT
    if "ОЦЕНКА" in rate_status:
        rate_cell.fill = INPUT_FILL
    ws.cell(row=row, column=5, value=rate_status)
    for c in range(1, 6):
        ws.cell(row=row, column=c).border = BORDER
        ws.cell(row=row, column=c).font = Font(name=FONT_NAME, size=9)
    row += 1
staff_last_row = row - 1

ws.cell(row=row, column=1, value="ИТОГО человек (лето)").font = Font(name=FONT_NAME, bold=True)
r_staff_summer_total = row
ws.cell(row=row, column=2, value=f"=SUM(B{staff_first_row}:B{staff_last_row})")
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value="ИТОГО человек (зима)").font = Font(name=FONT_NAME, bold=True)
ws.cell(row=row, column=2, value=f"=B{r_staff_summer_total}+SUM(C{staff_first_row}:C{staff_last_row})")
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 2

ws.cell(row=row, column=1, value="ФОТ/мес (лето)").font = Font(name=FONT_NAME, bold=True)
r_fot_summer = row
ws.cell(row=row, column=2, value=f"=SUMPRODUCT(B{staff_first_row}:B{staff_last_row},D{staff_first_row}:D{staff_last_row})").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value="ФОТ/мес (зима, включая доп.персонал)").font = Font(name=FONT_NAME, bold=True)
r_fot_winter = row
ws.cell(row=row, column=2, value=f"=B{r_fot_summer}+SUMPRODUCT(C{staff_first_row}:C{staff_last_row},D{staff_first_row}:D{staff_last_row})").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value="ФОТ-надбавка (30,2%, обязательна при официальном оформлении)")
r_fot_pct_note = row
ws.cell(row=row, column=2, value=f"=B{r_fot_summer}*0.302").number_format = CURRENCY_FMT
row += 2
for r in range(staff_last_row, row - 1):
    for c in (1, 2, 3, 4):
        ws.cell(row=r, column=c).border = BORDER

ws.cell(row=row, column=1, value="4. Расходные материалы (по нашим нормам consumable_norms.json)").font = SECTION_FONT
row += 1

params = ObjectParams(area_sqm=AREA_SQM, cleaning_days=26, general_days=1, staff_count=20, contract_months=1)
norms = load_consumable_norms()
_tz_items, _norm_qtys = [], []
for norm_id in norms:
    qty, unit, desc = estimate_consumable_quantity(norm_id, params)
    name = norms[norm_id]["display_name"]
    _tz_items.append(TZItem(raw_name=name, qty=qty, unit=unit))
    _norm_qtys.append((qty, unit, desc))

_price_lists = load_almin_price_list(discount_pct=5.0) + load_general_opt_price_list(supplier="ТК Сервис")
_match_results = resolve_all(_tz_items, _price_lists, supplier_adapters={"almin": stub_supplier_adapter})

from unit_price import estimate_purchase_quantity

headers_mat = ["Наименование материала", "Найдено в прайсе", "Поставщик", "Нужно (натур.)", "Упаковок", "Цена/уп, ₽", "Сумма/мес, ₽"]
for c, h in enumerate(headers_mat, start=1):
    ws.cell(row=row, column=c, value=h)
    ws.cell(row=row, column=c).fill = HEADER_FILL
    ws.cell(row=row, column=c).font = HEADER_FONT
    ws.cell(row=row, column=c).border = BORDER
row += 1
mat_first_row = row
for r, (qty, unit, desc) in zip(_match_results, _norm_qtys):
    ws.cell(row=row, column=1, value=r.tz_item.raw_name[:45])
    matched_name = r.matched_item.name[:30] if r.matched_item else "НЕ НАЙДЕНО"
    ws.cell(row=row, column=2, value=f"{matched_name} (conf {r.confidence:.2f})" if r.matched_item else matched_name)
    ws.cell(row=row, column=3, value=r.matched_item.supplier if r.matched_item else "—")
    ws.cell(row=row, column=4, value=f"{qty} {unit}")

    if r.matched_item and unit in ("л", "кг", "шт"):
        packs = estimate_purchase_quantity(qty, unit, r.matched_item)
    else:
        packs = None
    packs_cell = ws.cell(row=row, column=5, value=packs if packs is not None else qty)
    if packs is None:
        packs_cell.fill = INPUT_FILL

    price = r.matched_item.final_price if r.matched_item else 0
    price_cell = ws.cell(row=row, column=6, value=price)
    price_cell.number_format = CURRENCY_FMT
    if not r.matched_item:
        price_cell.fill = INPUT_FILL
    ws.cell(row=row, column=7, value=f"=E{row}*F{row}").number_format = CURRENCY_FMT
    for c in range(1, 8):
        ws.cell(row=row, column=c).border = BORDER
        if c != 1:
            ws.cell(row=row, column=c).font = Font(name=FONT_NAME, size=9)
    row += 1
mat_last_row = row - 1
ws.cell(row=row, column=1, value="ИТОГО материалы/мес").font = Font(name=FONT_NAME, bold=True)
r_mat_month = row
ws.cell(row=row, column=7, value=f"=SUM(G{mat_first_row}:G{mat_last_row})").number_format = CURRENCY_FMT
ws.cell(row=row, column=7).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1,
        value=f"Сопоставление с прайсами выполнено автоматически ({sum(1 for r in _match_results if r.matched_item)}/"
              f"{len(norms)} найдено). Количество — из consumable_norms.json (площадь×дни уборки); в самом "
              f"договоре списка материалов с количествами нет.").font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
row += 2

ws.cell(row=row, column=1, value="5. Оборудование (упомянуто в договоре — не входит в ФОТ)").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="Поломоечные машины (3 шт) + промышленные пылесосы (9 шт) — договор не указывает, "
                                    "аренда это или собственность Исполнителя. Наши расценки на технику "
                                    "(territory_equipment_rates.json) — про снегоуборку, сюда не подходят впрямую.")
ws.cell(row=row, column=1).font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
row += 2

ws.cell(row=row, column=1, value="6. Сверка: наша модель издержек vs реальная цена по договору").font = SECTION_FONT
row += 1
ws.cell(row=row, column=1, value="Наша оценка себестоимости/мес (лето): ФОТ + ФОТ-надбавка + материалы").font = Font(name=FONT_NAME, bold=True)
r_our_cost = row
ws.cell(row=row, column=2, value=f"=B{r_fot_summer}+B{r_fot_pct_note}+G{r_mat_month}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1, value="Реальная цена по договору/мес (лето)")
ws.cell(row=row, column=2, value=f"=B{r_price_summer}").number_format = CURRENCY_FMT
row += 1
ws.cell(row=row, column=1, value="Разница (накладные+прибыль Исполнителя+неучтённое)").font = Font(name=FONT_NAME, bold=True)
ws.cell(row=row, column=2, value=f"=B{r_price_summer}-B{r_our_cost}").number_format = CURRENCY_FMT
ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
row += 1
ws.cell(row=row, column=1,
        value="Если разница большая и положительная — наша модель, скорее всего, недооценивает издержки "
              "(неучтённое оборудование/аренда техники, накладные, налоговый режим 5% НДС) ИЛИ у Исполнителя "
              "высокая маржа. Если отрицательная — наши ОЦЕНОЧНЫЕ ставки завышены.").font = NOTE_FONT
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)

wb.save("smeta_krown_spetsekoservis.xlsx")
print("Смета сохранена: smeta_krown_spetsekoservis.xlsx")
