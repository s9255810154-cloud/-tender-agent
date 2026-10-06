"""
smeta_pipeline.py — универсальный конвейер «ТЗ → оценочная смета» для
НОВОГО тендера. Оценка по нормативам (knowledge_base), а не по реальным
зафиксированным данным уже выигранных контрактов — см. CLAUDE.md и
*_smeta_export.py (там реальные цифры конкретных договоров, здесь —
оценка для объекта, по которому ещё нет контракта).
"""
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook

from document_ingestion import ingest_document
from knowledge_base import ObjectParams, get_combined_shift_salary
from price_parser import load_general_opt_price_list, load_retail_suppliers_manual_additions
from smeta_builder import (
    CURRENCY_FMT, FINAL_FONT, PERCENT_FMT, TOTAL_FONT,
    SmetaSheet, build_contract_term_section, build_materials_table,
)
from tz_extraction import (
    detect_document_type, extract_tz_structured, extract_tz_structured_direct_area,
    process_extraction, process_premises_direct_area_extraction,
)


class PipelineError(Exception):
    """Понятная для UI ошибка на любом шаге конвейера."""


@dataclass
class SmetaPipelineResult:
    xlsx_path: str
    object_name: str
    region: str
    area_sqm: float
    cleaning_days: int
    contract_months: int
    staff_count: int
    staff_count_source: str
    fot_month: Optional[float]
    materials_unresolved: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _slugify(name: str) -> str:
    safe = re.sub(r"[^\w\-]+", "_", name, flags=re.UNICODE).strip("_")
    return safe[:60] or "smeta"


def _normalize_region(value: str) -> str:
    """«г. Москва» / «город Москва» / «Москва » -> «москва» — извлечённый
    Claude текст и выбор в форме иначе почти никогда не совпадут буквально."""
    normalized = value.strip().lower()
    normalized = re.sub(r"^(город\s+|г\.\s*|г\s+)", "", normalized)
    return normalized.strip(" .,")


def generate_smeta(
    file_path: str,
    *,
    region: str = "Москва",
    object_complexity: str = "стандартный",
    schedule_complexity: str = (
        "комбинированная смена (ежедневная-основная + "
        "ежедневная-поддерживающая, один сотрудник)"
    ),
    vat_rate: float = 0.20,
    output_dir: str = "Outputs/tender-agent-smeta",
    contract_months_override: Optional[int] = None,
) -> SmetaPipelineResult:
    warnings: list[str] = []

    # 1. Извлечение текста из документа
    doc = ingest_document(file_path)
    if len(doc.full_text.strip()) < 50:
        raise PipelineError(
            "Не удалось извлечь текст из документа (пусто или слишком мало "
            f"текста). Предупреждения: {'; '.join(doc.warnings) or 'нет'}."
        )

    # 2. Определить тип документа и выбрать схему извлечения
    doc_type = detect_document_type(doc)

    if doc_type == "premises_cleaning":
        extraction = extract_tz_structured(doc)
        if not extraction.get("contract_periods"):
            raise PipelineError(
                "В документе не найдены объёмы услуг по периодам (раздел "
                "«Перечень объектов закупки»). Автоматический расчёт "
                "невозможен — нужна ручная проверка документа."
            )
        summary = process_extraction(extraction)

    elif doc_type == "premises_cleaning_direct_area":
        extraction = extract_tz_structured_direct_area(doc)
        if not extraction.get("rooms"):
            raise PipelineError(
                "В документе не найдены строки с площадью и периодичностью "
                "по помещениям. Автоматический расчёт невозможен — нужна "
                "ручная проверка документа."
            )
        contract_months = contract_months_override
        if contract_months is None:
            contract_months = extraction.get("explicit_contract_months")
            if contract_months is not None:
                warnings.append(
                    f"Срок контракта определён из текста документа: {contract_months} мес. "
                    "Проверьте вручную перед использованием сметы."
                )
        if contract_months is None:
            raise PipelineError(
                "Для этого типа ТЗ срок контракта не задан в виде таблицы "
                "периодов и не найден явным числом в тексте — укажите срок "
                "контракта в месяцах в форме и повторите расчёт."
            )
        summary = process_premises_direct_area_extraction(extraction, contract_months)

    else:
        raise PipelineError(
            f"Тип ТЗ «{doc_type}» пока не поддерживается автоматическим "
            "расчётом сметы (поддерживаются только документы с прямой "
            "площадью по помещениям или с накопленным объёмом услуг за "
            "период). Нужна ручная проверка документа."
        )

    if doc_type == "premises_cleaning_direct_area":
        warnings.append(
            "«Дней уборки» для этого ТЗ — площадь-взвешенная оценка "
            "смешанной периодичности по помещениям, не точный факт из "
            "документа. Проверьте вручную перед использованием сметы."
        )
        if summary.cleaning_days == 0:
            warnings.append(
                "Не удалось определить периодичность уборки ни для одного "
                "помещения (текст периодичности не распознан) — база "
                "материалов в смете будет занижена, требуется ручная "
                "проверка документа и корректировка жёлтых ячеек."
            )

    if summary.staff_count_source == "норматив Роструда (площадь)":
        from knowledge_base import load_staff_norms
        scope_warning = load_staff_norms().get("_scope_warning")
        if scope_warning:
            warnings.append(f"Численность оценена по нормативу площади: {scope_warning}")

    if summary.region and _normalize_region(summary.region) != _normalize_region(region):
        warnings.append(
            f"В документе указан регион «{summary.region}», но расчёт выполнен "
            f"для региона «{region}» (выбран в форме) — нормативы ФОТ могут не "
            f"соответствовать фактическому региону объекта."
        )

    # 4. ФОТ — оценка по нормативу
    salary = get_combined_shift_salary(region, object_complexity, schedule_complexity)
    fot_month: Optional[float] = None
    if salary is not None:
        fot_month = salary["base_shift_salary"] * summary.staff_count
    else:
        warnings.append(
            f"Норматив ФОТ не найден для региона «{region}», сложности "
            f"«{object_complexity}», графика «{schedule_complexity}». "
            "Введите ФОТ вручную в сгенерированном файле (жёлтая ячейка) "
            "или дополните knowledge_base/salary_rules.json."
        )

    # 5. Сборка Excel
    wb = Workbook()
    sheet = SmetaSheet.new(
        wb, "Смета",
        {c: w for c, w in zip("ABCDEFGHI", [46, 16, 16, 8, 34, 22, 12, 14, 14])},
    )
    sheet.title_block(
        summary.object_name,
        subtitle=(
            "Автоматически сгенерировано (оценка по нормативам, НЕ "
            "финальная цена — накладные расходы и маржу добавляет эксперт)."
        ),
        region=region,
    )

    term_refs = build_contract_term_section(
        sheet, summary.contract_months, area_total_sqm=summary.area_sqm,
    )

    sheet.section("2. Численность и ФОТ (оценка по нормативу)")
    sheet.value_row(
        f"Численность персонала ({summary.staff_count_source})",
        summary.staff_count, editable=True,
    )
    if fot_month is not None:
        r_fot = sheet.value_row(
            "ФОТ, ₽/мес (норматив × численность)", fot_month,
            editable=True, fmt=CURRENCY_FMT,
        )
    else:
        r_fot = sheet.value_row(
            "ФОТ, ₽/мес — НОРМАТИВ НЕ НАЙДЕН, ввести вручную", 0,
            editable=True, fmt=CURRENCY_FMT,
        )
        sheet.note(
            f"Норматив не найден для «{region}» / «{object_complexity}» / "
            f"«{schedule_complexity}». Введите значение вручную в жёлтую "
            "ячейку выше."
        )
    sheet.blank()

    months = max(summary.contract_months, 1)
    params = ObjectParams(
        area_sqm=summary.area_sqm,
        cleaning_days=round(summary.cleaning_days / months),
        general_days=round(summary.general_days / months),
        staff_count=summary.staff_count,
        contract_months=1,
    )
    price_lists = (
        load_general_opt_price_list(supplier="ТК Сервис")
        + load_retail_suppliers_manual_additions()
    )
    mat = build_materials_table(sheet, params, price_lists, title="3. Расходные материалы")
    if mat.unresolved:
        warnings.append(
            f"Не найдено в прайсе (нужен ручной ввод цены): {', '.join(mat.unresolved)}"
        )

    sheet.section("4. ИТОГО — оценочная себестоимость")
    r_total_month = sheet.row
    sheet.ws.cell(row=r_total_month, column=1,
                  value="Себестоимость в месяц (ФОТ + материалы), без НДС").font = TOTAL_FONT
    sheet.ws.cell(row=r_total_month, column=2,
                  value=f"=B{r_fot}+{mat.total_cell}").number_format = CURRENCY_FMT
    sheet.ws.cell(row=r_total_month, column=2).font = TOTAL_FONT
    sheet.row += 1
    r_vat = sheet.value_row("Ставка НДС", vat_rate, editable=True, fmt=PERCENT_FMT)
    r_total_month_vat = sheet.row
    sheet.ws.cell(row=r_total_month_vat, column=1,
                  value="Себестоимость в месяц, с НДС").font = FINAL_FONT
    sheet.ws.cell(row=r_total_month_vat, column=2,
                  value=f"=B{r_total_month}*(1+B{r_vat})").number_format = CURRENCY_FMT
    sheet.ws.cell(row=r_total_month_vat, column=2).font = FINAL_FONT
    sheet.row += 1
    r_total_year = sheet.row
    sheet.ws.cell(row=r_total_year, column=1,
                  value="Себестоимость в год, с НДС").font = FINAL_FONT
    sheet.ws.cell(row=r_total_year, column=2,
                  value=f"=B{r_total_month_vat}*MIN(B{term_refs['months']},12)").number_format = CURRENCY_FMT
    sheet.ws.cell(row=r_total_year, column=2).font = FINAL_FONT
    sheet.row += 1
    sheet.note(
        "ОЦЕНОЧНАЯ себестоимость (ФОТ по нормативу + материалы по "
        "нормативу), НЕ финальная цена для тендера — накладные расходы и "
        "маржу эксперт добавляет отдельно (human-in-the-loop, см. CLAUDE.md)."
    )

    # 6. Сохранение
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    xlsx_path = out_dir / f"{_slugify(summary.object_name)}_{timestamp}.xlsx"
    wb.save(xlsx_path)

    return SmetaPipelineResult(
        xlsx_path=str(xlsx_path),
        object_name=summary.object_name,
        region=region,
        area_sqm=summary.area_sqm,
        cleaning_days=summary.cleaning_days,
        contract_months=summary.contract_months,
        staff_count=summary.staff_count,
        staff_count_source=summary.staff_count_source,
        fot_month=fot_month,
        materials_unresolved=mat.unresolved,
        warnings=warnings,
    )
