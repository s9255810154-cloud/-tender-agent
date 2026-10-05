"""
Document Ingestion (шаг 3 архитектуры) — единый вход для ТЗ в любом формате.

Раньше: PDF читался вручную, "слипшиеся" таблицы (Приложение 1 в реальном
ТЗ — площадь/дни уборки по месяцам) распутывались через ручной запуск
`pdftotext -layout` и разглядывание результата глазами. Здесь — то же самое,
но автоматически и для PDF/Word/Excel/сканов сразу.

Единый выход — IngestedDocument: список страниц/листов с текстом, отдельно
извлечённые таблицы (где получилось) и предупреждения о том, что могло
пойти не так (типичная ловушка — молча получить пустой/мусорный текст со
скана и не заметить). Дальше (шаг 4 — извлечение требований через LLM)
работает с этим объектом одинаково независимо от того, откуда он взялся.
"""
import re
import subprocess
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Optional


class DocFormat(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    IMAGE = "image"
    UNKNOWN = "unknown"


class ExtractionMethod(str, Enum):
    TEXT_LAYER = "text_layer"       # обычный текстовый слой PDF/Word
    TEXT_LAYER_LAYOUT = "text_layer_layout"  # pdftotext -layout — для таблиц, которые иначе "слипаются"
    TABLE_STRUCTURED = "table_structured"     # pdfplumber/openpyxl — таблица как таблица, не текст
    OCR = "ocr"                       # скан — понадобился tesseract


@dataclass
class ExtractedTable:
    page_or_sheet: str
    rows: list[list[str]]
    method: ExtractionMethod


@dataclass
class PageContent:
    index: int                       # номер страницы (PDF) или листа (Excel), с 1
    label: str                        # "Страница 3" / "Лист 'Данные'"
    text: str
    method: ExtractionMethod
    char_count: int = field(init=False)

    def __post_init__(self):
        self.char_count = len(self.text.strip())


@dataclass
class IngestedDocument:
    source_path: str
    doc_format: DocFormat
    pages: list[PageContent] = field(default_factory=list)
    tables: list[ExtractedTable] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    def summary(self) -> str:
        methods = {p.method.value for p in self.pages}
        return (f"{self.source_path}: {len(self.pages)} стр./листов, "
                f"методы={sorted(methods)}, таблиц={len(self.tables)}, "
                f"предупреждений={len(self.warnings)}")


# Порог "страница подозрительно пустая, похоже на скан без текстового слоя".
# Реальные текстовые страницы ТЗ почти всегда дают на 2+ порядка больше.
_EMPTY_PAGE_CHAR_THRESHOLD = 20


def detect_format(path: str) -> DocFormat:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return DocFormat.PDF
    if ext in (".docx", ".doc"):
        return DocFormat.DOCX
    if ext in (".xlsx", ".xls", ".xlsm"):
        return DocFormat.XLSX
    if ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
        return DocFormat.IMAGE
    return DocFormat.UNKNOWN


def ingest_document(path: str, ocr_lang: str = "rus+eng") -> IngestedDocument:
    """Единая точка входа. Дальше — диспетчер по формату."""
    fmt = detect_format(path)
    if fmt == DocFormat.PDF:
        return _ingest_pdf(path, ocr_lang)
    if fmt == DocFormat.DOCX:
        return _ingest_docx(path)
    if fmt == DocFormat.XLSX:
        return _ingest_excel(path)
    if fmt == DocFormat.IMAGE:
        return _ingest_image(path, ocr_lang)

    doc = IngestedDocument(source_path=path, doc_format=DocFormat.UNKNOWN)
    doc.warnings.append(f"Неизвестный формат файла: {Path(path).suffix}")
    return doc


# ============================== PDF ==============================

def _pdf_page_count(path: str) -> int:
    out = subprocess.run(["pdfinfo", path], capture_output=True, text=True, check=True).stdout
    m = re.search(r"^Pages:\s+(\d+)", out, re.MULTILINE)
    return int(m.group(1)) if m else 0


def _pdftotext_page(path: str, page: int, layout: bool) -> str:
    cmd = ["pdftotext"]
    if layout:
        cmd.append("-layout")
    cmd += ["-f", str(page), "-l", str(page), path, "-"]
    return subprocess.run(cmd, capture_output=True, text=True, check=True).stdout


def _ocr_pdf_page(path: str, page: int, ocr_lang: str) -> str:
    """Рендерит страницу в изображение и прогоняет через tesseract — для
    сканов без текстового слоя."""
    from pdf2image import convert_from_path
    images = convert_from_path(path, first_page=page, last_page=page, dpi=300)
    if not images:
        return ""
    import pytesseract
    return pytesseract.image_to_string(images[0], lang=ocr_lang)


def _extract_pdf_tables(path: str) -> list[ExtractedTable]:
    """pdfplumber даёт таблицу КАК ТАБЛИЦУ (список строк/ячеек), а не текст с
    потерянными разрывами колонок — там, где это получается, это надёжнее,
    чем полагаться на -layout и распутывать текст глазами постфактум."""
    tables: list[ExtractedTable] = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                for raw_table in page.extract_tables():
                    rows = [[cell or "" for cell in row] for row in raw_table]
                    if rows:
                        tables.append(ExtractedTable(
                            page_or_sheet=f"стр. {i}", rows=rows,
                            method=ExtractionMethod.TABLE_STRUCTURED,
                        ))
    except Exception:
        pass  # таблицы через pdfplumber — бонус, не критичный путь; текст всё равно извлечётся ниже
    return tables


def _ingest_pdf(path: str, ocr_lang: str) -> IngestedDocument:
    doc = IngestedDocument(source_path=path, doc_format=DocFormat.PDF)

    try:
        n_pages = _pdf_page_count(path)
    except Exception as e:
        doc.warnings.append(f"Не удалось прочитать метаданные PDF (pdfinfo): {e}")
        return doc

    for page in range(1, n_pages + 1):
        plain = _pdftotext_page(path, page, layout=False)
        layout = _pdftotext_page(path, page, layout=True)

        # Урок из ручной обработки реального ТЗ: обычный (не layout) режим
        # "слипает" многоколоночные таблицы в нечитаемую кашу. Берём -layout
        # как основной результат для страницы — он либо не хуже, либо
        # значительно лучше на табличных страницах.
        text = layout if layout.strip() else plain
        method = ExtractionMethod.TEXT_LAYER_LAYOUT if layout.strip() else ExtractionMethod.TEXT_LAYER

        if len(text.strip()) < _EMPTY_PAGE_CHAR_THRESHOLD:
            # Подозрительно пусто — похоже на скан без текстового слоя.
            try:
                ocr_text = _ocr_pdf_page(path, page, ocr_lang)
                if len(ocr_text.strip()) >= _EMPTY_PAGE_CHAR_THRESHOLD:
                    text, method = ocr_text, ExtractionMethod.OCR
                else:
                    doc.warnings.append(
                        f"Страница {page}: и текстовый слой, и OCR дали почти пустой "
                        f"результат — проверьте страницу вручную."
                    )
            except Exception as e:
                doc.warnings.append(f"Страница {page}: похоже на скан, но OCR не сработал ({e}).")

        doc.pages.append(PageContent(index=page, label=f"Страница {page}", text=text, method=method))

    doc.tables = _extract_pdf_tables(path)
    return doc


# ============================== DOCX ==============================

def _ingest_docx(path: str) -> IngestedDocument:
    """
    КРИТИЧНО: d.paragraphs (python-docx) не включает текст внутри таблиц —
    это отдельная структура документа. Реальные ТЗ почти всегда держат
    самое важное (объёмы по периодам, площади по помещениям, расходники)
    именно в таблицах Word — без явного рендера таблиц в текст Claude их
    просто не увидит (баг найден на реальном ТЗ, где вся площадь/
    периодичность уборки были в таблице, а full_text оставался пустым на
    это место). Поэтому таблицы рендерятся в текст СВЕРХ структурированного
    doc.tables, а не вместо него.
    """
    import docx
    doc = IngestedDocument(source_path=path, doc_format=DocFormat.DOCX)
    d = docx.Document(path)

    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    table_text_blocks = []
    for i, table in enumerate(d.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        doc.tables.append(ExtractedTable(
            page_or_sheet=f"таблица {i}", rows=rows, method=ExtractionMethod.TABLE_STRUCTURED,
        ))
        rendered_rows = [" | ".join(cell.strip() for cell in row) for row in rows if any(cell.strip() for cell in row)]
        if rendered_rows:
            table_text_blocks.append(f"[Таблица {i}]\n" + "\n".join(rendered_rows))

    if table_text_blocks:
        text = text + "\n\n" + "\n\n".join(table_text_blocks)

    doc.pages.append(PageContent(
        index=1, label="Документ целиком", text=text, method=ExtractionMethod.TEXT_LAYER,
    ))

    if not text.strip() and not doc.tables:
        doc.warnings.append("Документ Word не содержит ни текста, ни таблиц — проверьте файл вручную.")
    return doc


# ============================== XLSX ==============================

def _ingest_excel(path: str) -> IngestedDocument:
    import pandas as pd
    doc = IngestedDocument(source_path=path, doc_format=DocFormat.XLSX)

    xls = pd.ExcelFile(path)
    for i, sheet_name in enumerate(xls.sheet_names, start=1):
        df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
        rows = df.fillna("").astype(str).values.tolist()
        doc.tables.append(ExtractedTable(
            page_or_sheet=f"лист '{sheet_name}'", rows=rows, method=ExtractionMethod.TABLE_STRUCTURED,
        ))
        # Текстовое представление тоже даём — шагу 4 (LLM-извлечение) проще
        # работать с обычным текстом, чем с сырой таблицей значений.
        text = "\n".join("\t".join(row) for row in rows if any(cell.strip() for cell in row))
        doc.pages.append(PageContent(
            index=i, label=f"Лист '{sheet_name}'", text=text, method=ExtractionMethod.TABLE_STRUCTURED,
        ))

    if not doc.pages:
        doc.warnings.append("В книге Excel не найдено ни одного листа с данными.")
    return doc


# ============================== Изображение (скан без PDF-обёртки) ==============================

def _ingest_image(path: str, ocr_lang: str) -> IngestedDocument:
    import pytesseract
    from PIL import Image

    doc = IngestedDocument(source_path=path, doc_format=DocFormat.IMAGE)
    try:
        text = pytesseract.image_to_string(Image.open(path), lang=ocr_lang)
    except Exception as e:
        doc.warnings.append(f"OCR не сработал: {e}")
        text = ""

    doc.pages.append(PageContent(index=1, label="Изображение", text=text, method=ExtractionMethod.OCR))
    if len(text.strip()) < _EMPTY_PAGE_CHAR_THRESHOLD:
        doc.warnings.append("OCR дал почти пустой результат — проверьте изображение вручную (качество/поворот/язык).")
    return doc
